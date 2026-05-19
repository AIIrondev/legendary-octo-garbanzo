import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def generate_library_excel(items):
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"

    headers = [
        "Code", "Titel", "Autor", "Typ", "ISBN/Code", 
        "Filter 1", "Filter 2", "Filter 3", 
        "Status", "Ausgeliehen von", "Rückgabe", "Kosten"
    ]
    
    ws.append(headers)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for item in items:
        status = "Verfuegbar" if str(item.get("Verfuegbar", "True")).lower() == "true" else "Ausgeliehen"
        row = [
            item.get("Code_4", ""),
            item.get("Name", ""),
            item.get("Author", ""),
            item.get("ItemType", ""),
            item.get("ISBN", ""),
            item.get("Filter", ""),
            item.get("Filter2", ""),
            item.get("Filter3", ""),
            status,
            item.get("User", ""),
            item.get("ReturnDate", ""),
            item.get("Anschaffungskosten", "")
        ]
        ws.append(row)
        
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter 
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer
